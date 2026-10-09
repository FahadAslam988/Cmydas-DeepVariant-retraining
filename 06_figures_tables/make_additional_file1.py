"""Additional file 1 (PDF): Figure S1 (training curves and checkpoint selection) with title and legend."""
import os, subprocess
from docx import Document
from docx.shared import Cm, Pt
H = os.path.dirname(os.path.abspath(__file__)); P2 = os.path.join(H, '..', '..'); A = os.path.join(P2, '04_additional_files')
d = Document(); st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
p = d.add_paragraph(); r = p.add_run('Additional file 1'); r.bold = True
p = d.add_paragraph(); r = p.add_run('Figure S1. Training curves and checkpoint selection for the four retrained DeepVariant models.'); r.bold = True
d.add_picture(os.path.join(P2, '02_figures', 'FigS1_training_curves_checkpoint_selection.png'), width=Cm(16))
d.add_paragraph('Weighted F1 score on the training examples (grey) and on the validation contigs (blue; 15 contigs, 37.7 Mb, never used for training) '
    'during fine-tuning of DeepVariant v1.9.0 on simulated reads at a 70x, b 30x, c 15x and d 10x depth. The star and dotted line mark the '
    'selected checkpoint, the one with the highest validation F1 (DeepVariant\'s selection metric); its training step and validation F1 are given '
    'above each panel. At 15x and 10x, validation F1 declined after the selected checkpoint while training F1 continued to rise. The test contigs '
    'were not used for training or selection. These F1 values are DeepVariant\'s internal per-example training metric and are not comparable with the '
    'variant-level F1 scores from hap.py reported in the main text. The underlying values are given in Additional file 2 (sheet S5_training_curves).')
docx = os.path.join(A, 'Additional_file_1.docx'); d.save(docx)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', A, docx], capture_output=True, check=True)
os.remove(docx); print('saved', os.path.join(A, 'Additional_file_1.pdf'))
